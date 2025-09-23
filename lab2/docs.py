import pandas as pd
import os
import time
import subprocess
from docx import Document

items = []
total_sum = 0.0
i = 0

print("Имя Количество Цена (через пробел)")
while i < 5:
    line = input(f"Строка {i+1}: ").strip()
    parts = line.split()
    
    if len(parts) != 3:
        print("Ошибка: Имя, Количество, Цена")
        continue
    
    name = parts[0]
    try:
        quantity = int(parts[1])
        price = float(parts[2])
    except ValueError:
        print("Ошибка")
        continue
    
    amount = quantity * price
    total_sum += amount
    items.append({
        'Имя': name,
        'Количество': quantity,
        'Цена': price,
        'Сумма': amount
    })
    i += 1

df = pd.DataFrame(items)
df_sorted = df.sort_values(by='Цена', ascending=False).reset_index(drop=True)

file_excel = "отчет.xlsx"
total_row = {'Имя': 'Итого', 'Количество': '', 'Цена': '', 'Сумма': total_sum}
df_final = pd.concat([df_sorted, pd.DataFrame([total_row])], ignore_index=True)
df_final.to_excel(file_excel, index=False, sheet_name="Отчет")

file_docx = "отчет.docx"
doc = Document()
doc.add_paragraph('Список товаров:', style='Heading 2')
for _, row in df_sorted.iterrows():
    text = f"{row['Имя']} — {row['Количество']} шт × {row['Цена']:.2f} руб = {row['Сумма']:.2f} руб"
    doc.add_paragraph(text, style='List Bullet')
doc.add_paragraph('')
doc.add_paragraph(f'Общая сумма: {total_sum:.2f} руб.')
doc.save(file_docx)

file_pdf = "отчет.pdf"

from docx2pdf import convert
convert(file_docx, file_pdf)

abs_excel = os.path.abspath(file_excel)
abs_docx = os.path.abspath(file_docx)
abs_pdf = os.path.abspath(file_pdf)

print()
print(f"   Excel: {abs_excel}")
print(f"   Word:  {abs_docx}")
print(f"   PDF:   {abs_pdf}")

os.startfile(abs_excel)
os.startfile(abs_docx)
os.startfile(abs_pdf)

time.sleep(10)
        
subprocess.run(["taskkill", "/F", "/IM", "EXCEL.EXE"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
subprocess.run(["taskkill", "/F", "/IM", "FoxitPDFReader.exe"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

print("Все приложения закрыты.")
